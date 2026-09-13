import os

from pptx import Presentation

from app.schemas import (
    CFIUSInput,
    ClosingChecklistInput,
    JurisdictionTurnover,
    LBOInput,
    MasterDiligenceReport,
    RegulatoryHorizonInput,
)
from app.services.deal_financing import calculate_lbo
from app.services.pitchbook import create_pitchbook
from app.services.regulatory import cfius_screening, closing_checklist, regulatory_horizon


def _slide_titles(path: str) -> list[str]:
    prs = Presentation(path)
    titles = []
    for slide in prs.slides:
        if slide.shapes.title is not None:
            titles.append(slide.shapes.title.text)
    return titles


def _base_report(**extra) -> MasterDiligenceReport:
    return MasterDiligenceReport(
        deal_id="pitchbook-test",
        company_b_financials={"latest_revenue": 42_000_000},
        valuation={"enterprise_value": 87_500_000},
        synergies={"financial_synergies": {"total_annual_synergy": 5_200_000}},
        **extra,
    )


def test_pitchbook_without_m6_m7_has_four_slides(tmp_path):
    os.environ["PITCHBOOK_OUTPUT_DIR"] = str(tmp_path)
    path = create_pitchbook(_base_report())
    titles = _slide_titles(path)
    assert len(titles) == 4
    assert "Financing & Returns" not in titles
    assert "Regulatory & Closing" not in titles


def test_pitchbook_includes_m6_m7_slides(tmp_path):
    os.environ["PITCHBOOK_OUTPUT_DIR"] = str(tmp_path)

    financing = calculate_lbo(LBOInput(enterprise_value=500.0, ebitda=100.0, exit_multiple=6.0))
    regulatory = {
        "regulatory_horizon": regulatory_horizon(RegulatoryHorizonInput(
            deal_ev=2_000_000_000.0, combined_market_share=0.35, product_overlaps=4, sector="telecom",
            jurisdictions=[JurisdictionTurnover(jurisdiction="US_HSR", combined_turnover_usd=500_000_000.0)],
        )),
        "cfius": cfius_screening(CFIUSInput(us_business=True, foreign_acquirer=True, acquirer_country_of_concern=True, critical_technology=True, sensitive_personal_data=True)),
        "closing_checklist": closing_checklist(ClosingChecklistInput(required_jurisdictions=["US_HSR"], change_of_control_consents=2)),
    }

    report = _base_report(financing=financing, regulatory=regulatory)
    path = create_pitchbook(report)
    titles = _slide_titles(path)

    assert len(titles) == 6
    assert "Financing & Returns" in titles
    assert "Regulatory & Closing" in titles
