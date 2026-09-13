from __future__ import annotations

import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.schemas import MasterDiligenceReport


class PitchbookGenerationError(ValueError):
    pass


class PitchbookDependencyError(PitchbookGenerationError):
    pass


def _import_pptx() -> Any:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise PitchbookDependencyError(
            "python-pptx is not installed. Install with `pip install python-pptx`."
        ) from exc
    return Presentation


def _safe_float(value: Any) -> float | None:
    try:
        return float(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _currency(value: float | None) -> str:
    return "N/A" if value is None else f"${value:,.2f}"


def generate_pitchbook(master_json: MasterDiligenceReport | dict[str, Any]) -> str:
    report = (
        master_json
        if isinstance(master_json, MasterDiligenceReport)
        else MasterDiligenceReport.model_validate(master_json)
    )

    Presentation = _import_pptx()
    prs = Presentation()

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Project Alpha: M&A Advisory Overview"
    subtitle = slide.placeholders[1]
    subtitle.text = (
        f"Deal ID: {report.deal_id}\n"
        "AI-Generated Deal Strategy & Diligence\n"
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )

    # Slide 2
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Financial Valuation & Synergies"
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    target_revenue = _safe_float(report.company_b_financials.get("latest_revenue"))
    total_synergy = _safe_float(
        report.synergies.get("financial_synergies", {}).get("total_annual_synergy")
    )
    enterprise_value = _safe_float((report.valuation or {}).get("enterprise_value"))

    tf.paragraphs[0].text = f"Target Enterprise Value: {_currency(enterprise_value)}"
    p = tf.add_paragraph()
    p.text = f"Target Base Revenue: {_currency(target_revenue)}"
    p = tf.add_paragraph()
    p.text = f"Estimated Annual Synergies: {_currency(total_synergy)}"

    # Slide 3
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Due Diligence: Key Legal Risks"
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    risks = ((report.legal_risks or {}).get("risks") or [])[:6]
    if risks:
        tf.paragraphs[0].text = "AI Scanner flagged the following:"
        for risk in risks:
            p = tf.add_paragraph()
            p.text = (
                f"[{risk.get('severity', 'Unknown')}] {risk.get('risk_category', 'Risk')}: "
                f"{risk.get('ai_explanation', '')}"
            )
            p.level = 1
    else:
        tf.paragraphs[0].text = "No major legal risks identified or scan bypassed."

    # Slide 4
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Strategic Recommendation & Tactics"
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    deal_structure = (report.deal_structure or {}).get("recommended_structure", "N/A")
    opening_bid = _safe_float((report.negotiation_strategy or {}).get("opening_bid_price"))
    walk_away = _safe_float((report.negotiation_strategy or {}).get("walk_away_price"))
    tactics = (report.negotiation_strategy or {}).get("tactical_moves", [])

    tf.paragraphs[0].text = f"Recommended Structure: {deal_structure}"
    p = tf.add_paragraph()
    p.text = f"Suggested Opening Bid: {_currency(opening_bid)}"
    p = tf.add_paragraph()
    p.text = f"Hard Walk-Away Price: {_currency(walk_away)}"

    if tactics:
        p = tf.add_paragraph()
        p.text = "Tactics:"
        for tactic in tactics[:5]:
            p = tf.add_paragraph()
            p.text = str(tactic)
            p.level = 1
    else:
        p = tf.add_paragraph()
        p.text = "Strategy generation skipped or unavailable."

    # Slide 5 — Financing & Returns (M6), only when financing data is present.
    financing = report.financing or {}
    if financing:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Financing & Returns"
        tf = slide.shapes.placeholders[1].text_frame
        tf.clear()

        returns = financing.get("returns", {})
        su = financing.get("sources_and_uses", {})
        covenants = financing.get("covenants", {})
        irr = _safe_float(returns.get("equity_irr"))
        moic = _safe_float(returns.get("moic"))
        entry_x = _safe_float(financing.get("entry_multiple"))
        exit_x = _safe_float(financing.get("exit_multiple"))
        min_dscr = _safe_float(covenants.get("min_dscr"))

        tf.paragraphs[0].text = (
            f"Entry / Exit Multiple: {entry_x:.1f}x / {exit_x:.1f}x"
            if entry_x is not None and exit_x is not None
            else "Entry / Exit Multiple: N/A"
        )
        lines = [
            f"Total Debt Raised: {_currency(_safe_float(financing.get('total_debt')))}",
            f"Sponsor Equity: {_currency(_safe_float(financing.get('equity_contribution')))}",
            f"Equity IRR: {irr:.1%}" if irr is not None else "Equity IRR: N/A",
            f"MOIC: {moic:.2f}x" if moic is not None else "MOIC: N/A",
            f"Meets Board Hurdle: {'Yes' if returns.get('meets_hurdle') else 'No'}",
            f"Minimum DSCR: {min_dscr:.2f}x" if min_dscr is not None else "Minimum DSCR: N/A",
        ]
        for line in lines:
            p = tf.add_paragraph()
            p.text = line

        stack = financing.get("capital_stack") or []
        if stack:
            p = tf.add_paragraph()
            p.text = "Capital Stack:"
            for tranche in stack:
                amount = _safe_float(tranche.get("amount"))
                if not amount:
                    continue
                pct = _safe_float(tranche.get("pct_of_capital"))
                pct_text = f" ({pct:.0%})" if pct is not None else ""
                p = tf.add_paragraph()
                p.text = f"{tranche.get('tranche', 'Tranche')}: {_currency(amount)}{pct_text}"
                p.level = 1

    # Slide 6 — Regulatory & Closing (M7), only when regulatory data is present.
    regulatory = report.regulatory or {}
    if regulatory:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Regulatory & Closing"
        tf = slide.shapes.placeholders[1].text_frame
        tf.clear()
        first_line_set = False

        def _line(text: str, level: int = 0) -> None:
            nonlocal first_line_set
            if not first_line_set:
                tf.paragraphs[0].text = text
                tf.paragraphs[0].level = level
                first_line_set = True
            else:
                p = tf.add_paragraph()
                p.text = text
                p.level = level

        horizon = regulatory.get("regulatory_horizon") or {}
        if horizon:
            _line(f"Antitrust filings required: {horizon.get('filings_required', 'N/A')}")
            p2 = _safe_float(horizon.get("phase2_probability_model"))
            if p2 is not None:
                _line(f"Phase II / second-request probability: {p2:.0%}", 1)
            crit = horizon.get("critical_path")
            if crit:
                _line(f"Critical path: {crit.get('jurisdiction')} ({crit.get('expected_review_weeks')} weeks)", 1)
            divest = _safe_float(horizon.get("total_estimated_divestiture_cost_usd"))
            if divest is not None:
                _line(f"Estimated divestiture cost: {_currency(divest)}", 1)

        cfius = regulatory.get("cfius") or {}
        if cfius:
            _line(f"CFIUS risk tier: {cfius.get('risk_tier', 'N/A')}"
                  + (" (blocking)" if cfius.get("blocking") else ""))
            timeline = cfius.get("estimated_timeline_days")
            if timeline:
                _line(f"CFIUS review timeline: {timeline} days", 1)

        checklist = regulatory.get("closing_checklist") or {}
        if checklist:
            _line(f"Closing checklist items: {checklist.get('total_items', 'N/A')}")

        if not first_line_set:
            _line("Regulatory analysis not available.")

    out_dir = Path(os.getenv("PITCHBOOK_OUTPUT_DIR", "generated_pitchbooks"))
    out_dir.mkdir(parents=True, exist_ok=True)
    filepath = out_dir / f"MA_Pitchbook_{report.deal_id}.pptx"
    prs.save(str(filepath))
    return str(filepath.resolve())


def create_pitchbook(report: MasterDiligenceReport | dict[str, Any]) -> str:
    return generate_pitchbook(report)

